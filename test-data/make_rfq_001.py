"""Generate synthetic test RFQ files (PDF + Excel + answer key) for the quote agent.

All companies and people are fictional. SKUs and prices are real Petro Valve
prod catalog values as of 2026-10-16, so the files can also score parsing and
SKU matching later.

Run: python make_rfq_001.py  (needs reportlab and openpyxl)
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

OUT = Path(__file__).parent / "rfq-001"
OUT.mkdir(exist_ok=True)

CUSTOMER = {
    "company": "Bayou Process Supply (TEST)",
    "contact": "Dana Whitfield (fictional)",
    "address": "4100 Example Pkwy, Pasadena, TX 77503",
    "rfq_no": "BPS-RFQ-TEST-001",
}

# expect: quote | recommend (best guess, disclosed) | escalate:<reason>
PDF_LINES = [
    ("VO 12111-1", 4, "exact SKU", "quote", "VO 12111-1", 135.00),
    ('Vogt 12111 1-1/2"', 2, "partial SKU", "quote", "VO 12111-1-1/2", 283.00),
    ('Gate valve 3/4", 800#, A105 forged, socket weld, Vogt', 6, "description only", "recommend", "VO SW12141-3/4 or VO 12111-3/4", None),
    ("CR 147XU-2", 1, "exact SKU", "quote", "CR 147XU-2", 2674.39),
    ('VO 12111-3', 2, "exact SKU, $0.01 placeholder price", "escalate:price_anomaly", "VO 12111-3", 0.01),
    ('Fisher ED control valve 4" CL300', 1, "not in catalog", "escalate:non_catalog", None, None),
]

XLSX_LINES = [
    ("VO 12111-1/2", "", 10, "quote", "VO 12111-1/2", 98.00),
    ("VO 12111-3/4", "", 10, "quote", "VO 12111-3/4", 110.00),
    ("SW12141", '1"', 5, "quote", "VO SW12141-1", 237.00),
    ("VO SW12141-2", "", 3, "quote", "VO SW12141-2", 657.00),
    ("VO 12111F8M-1", "", 4, "quote", "VO 12111F8M-1", 305.00),
    ("12111F8M", "2 inch", 2, "quote", "VO 12111F8M-2", 988.00),
    ("CR 147LUN-4", "", 1, "quote", "CR 147LUN-4", 10792.82),
    ("Crane 147XU", '3"', 1, "quote", "CR 147XU-3", 3398.97),
    ("VO SW12141-1-1/4", "", 2, "escalate:price_anomaly", "VO SW12141-1-1/4", 0.01),
    ("VO 12111-1", "", 25000, "escalate:unusual_quantity", "VO 12111-1", 135.00),
]

EMAIL_BODY = f"""Hi,

Please quote the items on the attached RFQ ({CUSTOMER['rfq_no']}) and the attached spreadsheet.
Also add 3 ea of Vogt 12111 1/2" gate valves to the quote.

Ship to our Pasadena, TX yard. Need pricing and lead time by Friday.

Thanks,
Dana Whitfield
Purchasing, Bayou Process Supply
4100 Example Pkwy, Pasadena, TX 77503

[TEST EMAIL - synthetic RFQ for Petro Valve quote agent testing. Company and contact are fictional.]
"""
BODY_LINES = [('Vogt 12111 1/2" gate valve', 3, "quote", "VO 12111-1/2", 98.00)]


def make_pdf(path):
    c = canvas.Canvas(str(path), pagesize=letter)
    w, h = letter
    y = h - 0.9 * inch
    c.setFont("Helvetica-Bold", 16)
    c.drawString(0.8 * inch, y, "REQUEST FOR QUOTATION")
    c.setFont("Helvetica", 9)
    c.drawRightString(w - 0.8 * inch, y, "TEST DOCUMENT - synthetic data")
    y -= 0.35 * inch
    c.setFont("Helvetica", 10)
    for label, val in [("From", CUSTOMER["company"]), ("Contact", CUSTOMER["contact"]),
                       ("Address", CUSTOMER["address"]), ("RFQ No.", CUSTOMER["rfq_no"]),
                       ("Date", "10/16/2026"), ("Quote due", "10/23/2026")]:
        c.drawString(0.8 * inch, y, f"{label}:")
        c.drawString(1.7 * inch, y, val)
        y -= 0.2 * inch
    y -= 0.2 * inch
    c.setFont("Helvetica-Bold", 10)
    c.drawString(0.8 * inch, y, "Line")
    c.drawString(1.3 * inch, y, "Description / Part No.")
    c.drawRightString(6.6 * inch, y, "Qty")
    c.drawRightString(7.6 * inch, y, "UOM")
    y -= 0.08 * inch
    c.line(0.8 * inch, y, w - 0.8 * inch, y)
    c.setFont("Helvetica", 10)
    for i, (desc, qty, *_rest) in enumerate(PDF_LINES, 1):
        y -= 0.25 * inch
        c.drawString(0.8 * inch, y, str(i))
        c.drawString(1.3 * inch, y, desc)
        c.drawRightString(6.6 * inch, y, str(qty))
        c.drawRightString(7.6 * inch, y, "EA")
    y -= 0.5 * inch
    c.setFont("Helvetica", 9)
    for note in ["Notes: Quote FOB Houston. Include lead time per line.",
                 "Substitutions must be noted. MTRs required for carbon steel."]:
        c.drawString(0.8 * inch, y, note)
        y -= 0.18 * inch
    c.save()


def make_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "RFQ"
    ws.append([f"{CUSTOMER['company']} - additional items for {CUSTOMER['rfq_no']}  (TEST DATA)"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Item", "Part #", "Size", "Qty", "Notes"])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for i, (part, size, qty, *_rest) in enumerate(XLSX_LINES, 1):
        ws.append([i, part, size, qty, ""])
    for col, width in zip("ABCDE", (6, 20, 10, 8, 24)):
        ws.column_dimensions[col].width = width
    wb.save(path)


def make_answer_key(path):
    key = {"rfq": CUSTOMER["rfq_no"], "source_prices": "petrovalve.com prod, 2026-10-16", "lines": []}
    for desc, qty, kind, expect, sku, price in PDF_LINES:
        key["lines"].append({"source": "pdf", "input": desc, "qty": qty, "case": kind,
                             "expect": expect, "sku": sku, "unit_price": price})
    for part, size, qty, expect, sku, price in XLSX_LINES:
        key["lines"].append({"source": "xlsx", "input": f"{part} {size}".strip(), "qty": qty,
                             "expect": expect, "sku": sku, "unit_price": price})
    for desc, qty, expect, sku, price in BODY_LINES:
        key["lines"].append({"source": "email_body", "input": desc, "qty": qty,
                             "expect": expect, "sku": sku, "unit_price": price})
    path.write_text(json.dumps(key, indent=2), encoding="utf-8")


make_pdf(OUT / "BPS-RFQ-TEST-001.pdf")
make_xlsx(OUT / "BPS-RFQ-TEST-001-additional-items.xlsx")
make_answer_key(OUT / "expected.json")
(OUT / "email_body.txt").write_text(EMAIL_BODY, encoding="utf-8")
print("wrote", sorted(p.name for p in OUT.iterdir()))
